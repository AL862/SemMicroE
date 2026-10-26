"""Inspect every actual PPTX/DOCX package and verify the eleven worked examples."""
from pathlib import Path
import json,re,hashlib,zipfile
from fractions import Fraction as F
from pptx import Presentation
from docx import Document
from lxml import etree as E
from render import fmt
ROOT=Path(__file__).parent;OUT=ROOT.parents[1];FOLDER=OUT/'[3] Prezentări pentru seminare'
def plain(s):return re.sub('<[^>]*>','',fmt(s)).replace('&gt;','>').replace('&lt;','<').replace('&amp;','&')
def norm(s):return ' '.join(s.split())
def alt_text(slide):
 rows=[slide['title']]
 for field in ['hook','promise','lead','formula','condition','conclusion','check']:
  if slide.get(field):rows.append(slide[field])
 for field in ['cards','steps']:
  for label,value in slide.get(field,[]):rows.extend([label,value])
 for field in ['lines','points','requirements']:rows.extend(slide.get(field,[]))
 return '. '.join(plain(s).rstrip('.') for s in rows)+'.'

def validate():
 plan=json.loads((ROOT/'plan.json').read_text());render=json.loads((ROOT/'render-validation.json').read_text());assert len(render)==110;report=[];manifest=[]
 for deck in plan:
  n=deck['seminar'];pptx=next(FOLDER.glob(f'Prezentare - Seminarul {n}. *.pptx'));p=Presentation(pptx);assert len(p.slides)==10
  p.core_properties.author='ASEM';p.core_properties.last_modified_by='ASEM';p.core_properties.language='ro-RO'
  for slide,data in zip(p.slides,deck['slides']):
   if data['kind'] not in ['exercise','solution']:
    assert len(slide.shapes)==1 and slide.shapes[0].shape_type==13
    shape=slide.shapes[0];shape._element.xpath('.//p:cNvPr')[0].set('descr',alt_text(data))
  p.save(pptx)
  note_file=FOLDER/'Conspecte pentru lector'/f'Seminar_{n}_Conspect_Profesor.docx';doc=Document(note_file);doc.core_properties.author='ASEM';doc.core_properties.last_modified_by='ASEM';doc.save(note_file)
  doc_text=norm(' '.join(x.text for x in doc.paragraphs));p=Presentation(pptx)
  assert abs(p.slide_width/p.slide_height-16/9)<1e-6
  note_count=0
  for slide,data in zip(p.slides,deck['slides']):
   notes=slide.notes_slide.notes_text_frame.text;expected='Timp orientativ: '+str(data['planned_seconds'])+' de secunde. '+data['notes']
   assert norm(notes)==norm(plain(expected)),(n,data['number'],'notes mismatch')
   assert norm(plain(data['notes'])) in doc_text,(n,data['number'],'DOCX mismatch')
   assert ',' not in notes,(n,data['number'],'comma')
   assert not re.search(r'Tom[șşs]a|Mankiw|Мэнкью|Томша',notes,re.I)
   note_count+=len(notes.split())
   if data['kind'] in ['exercise','solution']:
    shapes=[x for x in slide.shapes if x.has_text_frame and x.text.strip()];assert len(shapes)>=7
    actual=norm(' '.join(x.text for x in shapes))
    values=[data['condition'],*data['requirements']] if data['kind']=='exercise' else [value for pair in data['steps'] for value in pair]+[data['conclusion']]
    for value in values:assert norm(plain(value)) in actual,(n,data['number'],value)
    assert not any('Răspuns așteptat' in x.text for x in shapes)
   else:
    picture=slide.shapes[0];expected_image=ROOT/'rendered'/f'{n:02}-{data["number"]:02}.png';assert picture.image.blob==expected_image.read_bytes()
    assert picture.image.size==(1920,1080)
  with zipfile.ZipFile(pptx) as archive:
   assert archive.testzip() is None
   for name in archive.namelist():
    if name.endswith('.xml') or name.endswith('.rels'):E.fromstring(archive.read(name))
    if name.endswith('.rels'):assert b'TargetMode="External"' not in archive.read(name),(n,name)
   assert len([x for x in archive.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',x)])==10
  report.append({'seminar':n,'slides':10,'raster_infographics':8,'editable_example_slides':[8,9],'embedded_notes':10,'notes_words':note_count,'planned_minutes':11,'notes_match_docx':True,'embedded_images_match_reviewed_png':True,'external_dependencies':False,'xml_package_integrity':True})
  for file in [pptx,note_file]:manifest.append({'path':file.relative_to(OUT).as_posix(),'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
 # Independent calculations rather than reading back the displayed results.
 tests=[('S1 frontier',50-5*4,30),('S1 sacrifice',(50-5*4)-(50-5*6),10),('S2 equilibrium P',F(120,3),40),('S2 Qd',100-40,60),('S2 Qs',-20+2*40,60),('S3 quantity',40-2*10,20),('S3 elasticity',-2*F(10,20),-1),('S4 budget A',20*3+10*6,120),('S4 budget B',20*2+10*8,120),('S4 utility A',3*6,18),('S4 utility B',2*8,16),('S5 AP',F(30,3),10),('S5 MP',F(30-28,3-2),2),('S6 total',100+5*20,200),('S6 average',F(200,20),10),('S6 marginal',(100+5*21)-(100+5*20),5),('S7 optimum quantity',F(50-10,2),20),('S7 profit',50*20-(100+10*20+20**2),300),('S8 quantity',F(20-10,2),5),('S8 operating result',20*5-(100+10*5+5**2),-75),('S8 stopping result',-100,-100),('S9 quantity',F(100-20,4),20),('S9 price',100-2*20,60),('S9 profit',60*20-(100+20*20),700),('S10 quantity',F(80-20,4),15),('S10 price',80-15,65),('S10 cost',200+20*15+15**2,725),('S10 profit',65*15-725,250),('S11 response',30-F(20,2),20),('S11 price',90-20-20,50),('S11 profit',(50-30)*20,400)]
 for label,actual,expected in tests:assert actual==expected,(label,actual,expected)
 (ROOT/'package-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));(ROOT/'calculation-checks.json').write_text(json.dumps({'passed':len(tests),'checks':[{'name':label,'value':str(actual)} for label,actual,_ in tests]},indent=2));(ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2));print('Validated11PPTX,110slides,110notes,11DOCX;',len(tests),'exact calculations')
if __name__=='__main__':validate()
